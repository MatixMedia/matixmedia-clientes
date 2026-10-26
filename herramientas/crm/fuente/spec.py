"""Convierte MATIXMEDIA_CRM.xlsx (build_crm.py) en una especificación JSON compacta
que el script de Apps Script usa para construir la hoja (Instalar CRM)."""
import json, sys, datetime as dt
import openpyxl
from openpyxl.utils import get_column_letter

wb = openpyxl.load_workbook(sys.argv[1])
spec = []
for ws in wb.worksheets:
    s = {"n": ws.title, "tab": (ws.sheet_properties.tabColor.rgb[-6:] if ws.sheet_properties.tabColor is not None else None),
         "frz": ws.freeze_panes, "w": {}, "c": [], "dv": []}
    for k, d in ws.column_dimensions.items():
        if d.width and abs(d.width - 13) > 0.01:
            s["w"][k] = round(d.width)
    for row in ws.iter_rows():
        for c in row:
            if c.value is None and not (c.fill and c.fill.fgColor and c.fill.fill_type == "solid"):
                continue
            e = {"a": c.coordinate}
            v = c.value
            if isinstance(v, (dt.datetime, dt.date)):
                v = v.isoformat()
            if v is not None:
                e["v"] = v
            if c.font is not None:
                if c.font.bold: e["b"] = 1
                col = c.font.color.rgb if c.font.color is not None and isinstance(c.font.color.rgb, str) else None
                if col and col[-6:] not in ("000000",): e["fc"] = col[-6:]
                if c.font.size and c.font.size != 10: e["fs"] = int(c.font.size)
            if c.fill is not None and c.fill.fill_type == "solid":
                e["bg"] = c.fill.fgColor.rgb[-6:]
            if c.number_format and c.number_format != "General":
                e["nf"] = c.number_format
            if c.alignment is not None and c.alignment.wrap_text:
                e["wr"] = 1
            if len(e) > 1:
                s["c"].append(e)
    for dv in ws.data_validations.dataValidation:
        s["dv"].append({"r": str(dv.sqref), "f": dv.formula1})
    spec.append(s)
out = json.dumps(spec, ensure_ascii=False, separators=(",", ":"))
open(sys.argv[2], "w").write(out)
print(len(out), "bytes", [s["n"] for s in spec], sum(len(s["c"]) for s in spec), "celdas")
